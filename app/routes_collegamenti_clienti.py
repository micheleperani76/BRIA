#!/usr/bin/env python3
# ==============================================================================
# GESTIONE FLOTTA - Route Collegamenti Clienti
# ==============================================================================
# Versione: 2.0.0
# Data: 2026-01-26
# Descrizione: Gestione collegamenti/relazioni tra clienti
# Novita' v2: Supporto relazioni inverse (CODICE_INVERSO da Excel)
# ==============================================================================

from flask import Blueprint, request, jsonify, render_template, session
from .database import get_connection
from .auth import login_required, permesso_richiesto
from .database_utenti import get_subordinati
from datetime import datetime

collegamenti_bp = Blueprint('collegamenti', __name__)

# ==============================================================================
# FUNZIONE: CARICA TIPI RELAZIONE DA EXCEL (con CODICE_INVERSO)
# ==============================================================================

def get_tipi_relazione():
    """Carica i tipi di relazione dal file Excel (con codice inverso)."""
    import openpyxl
    import os
    
    filepath = os.path.join(os.path.dirname(__file__), '..', 'impostazioni', 'tipi_relazione.xlsx')
    
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        ws = wb.active
        
        tipi = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                tipi.append({
                    'codice': row[0],
                    'descrizione': row[1],
                    'codice_inverso': row[2] if len(row) > 2 and row[2] else row[0]  # Fallback: stesso codice
                })
        
        wb.close()
        return tipi
    except Exception as e:
        print(f"Errore lettura tipi_relazione.xlsx: {e}")
        # Fallback
        return [
            {'codice': 'COL', 'descrizione': 'Collegato', 'codice_inverso': 'COL'},
            {'codice': 'CONS', 'descrizione': 'Consociata', 'codice_inverso': 'CONS'},
            {'codice': 'FIL', 'descrizione': 'Filiale', 'codice_inverso': 'SEDE'},
        ]


def get_descrizione_relazione(codice):
    """Ritorna la descrizione per un codice relazione."""
    tipi = get_tipi_relazione()
    for t in tipi:
        if t['codice'] == codice:
            return t['descrizione']
    return codice  # Fallback: ritorna il codice stesso


def get_codice_inverso(codice):
    """Ritorna il codice inverso per una relazione."""
    tipi = get_tipi_relazione()
    for t in tipi:
        if t['codice'] == codice:
            return t.get('codice_inverso', codice)
    return codice  # Fallback: stesso codice


def get_descrizione_per_vista(codice, usa_inverso=False):
    """
    Ritorna la descrizione corretta in base al punto di vista.
    
    Se usa_inverso=True, cerca il codice inverso e ne restituisce la descrizione.
    Esempio: codice='CLI', usa_inverso=True -> cerca 'FORN' -> 'Fornitore'
    """
    tipi = get_tipi_relazione()
    
    if usa_inverso:
        # Prima trovo il codice inverso
        codice_da_cercare = codice
        for t in tipi:
            if t['codice'] == codice:
                codice_da_cercare = t.get('codice_inverso', codice)
                break
        
        # Poi trovo la descrizione del codice inverso
        for t in tipi:
            if t['codice'] == codice_da_cercare:
                return t['descrizione']
    else:
        # Descrizione diretta
        for t in tipi:
            if t['codice'] == codice:
                return t['descrizione']
    
    return codice  # Fallback


# ==============================================================================
# FUNZIONI HELPER
# ==============================================================================

def get_collegamenti_cliente(conn, cliente_id):
    """
    Recupera tutti i collegamenti attivi di un cliente.
    
    Aggiunge il campo 'descrizione_relazione' con la descrizione corretta
    in base al punto di vista (usa l'inverso se il cliente e' cliente_b).
    """
    cursor = conn.cursor()
    cursor.execute('''
        SELECT 
            cc.id,
            cc.tipo_relazione,
            cc.note,
            cc.creato_il,
            cc.creato_da,
            cc.cliente_a_id,
            cc.cliente_b_id,
            u.nome || ' ' || u.cognome as creato_da_nome,
            CASE 
                WHEN cc.cliente_a_id = ? THEN cc.cliente_b_id 
                ELSE cc.cliente_a_id 
            END as altro_cliente_id,
            c.ragione_sociale as altro_cliente_nome,
            c.p_iva as altro_cliente_piva,
            c.cod_fiscale as altro_cliente_cf,
            c.commerciale_id
        FROM collegamenti_clienti cc
        JOIN clienti c ON c.id = CASE 
            WHEN cc.cliente_a_id = ? THEN cc.cliente_b_id 
            ELSE cc.cliente_a_id 
        END
        LEFT JOIN utenti u ON u.id = cc.creato_da
        WHERE (cc.cliente_a_id = ? OR cc.cliente_b_id = ?)
        AND cc.attivo = 1
        ORDER BY cc.creato_il DESC
    ''', (cliente_id, cliente_id, cliente_id, cliente_id))
    
    columns = [desc[0] for desc in cursor.description]
    risultati = []
    
    for row in cursor.fetchall():
        collegamento = dict(zip(columns, row))
        
        # Determina se usare la relazione inversa
        # Se il cliente corrente e' cliente_b, devo mostrare la relazione inversa
        usa_inverso = (collegamento['cliente_b_id'] == cliente_id)
        
        # Aggiungi la descrizione corretta per la vista
        collegamento['descrizione_relazione'] = get_descrizione_per_vista(
            collegamento['tipo_relazione'], 
            usa_inverso=usa_inverso
        )
        
        # Flag utile per il frontend
        collegamento['vista_inversa'] = usa_inverso
        
        risultati.append(collegamento)
    
    return risultati


def get_identificativo_cliente(cliente):
    """Restituisce l'identificativo fiscale del cliente (P.IVA o CF)."""
    piva = cliente.get('altro_cliente_piva') or cliente.get('p_iva')
    cf = cliente.get('altro_cliente_cf') or cliente.get('cod_fiscale')
    
    if piva and piva.strip():
        return f"IT{piva.strip()}"
    elif cf and cf.strip():
        return cf.strip()
    return None


# ==============================================================================
# ROUTE: API TIPI RELAZIONE
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/tipi-relazione')
@login_required
def api_tipi_relazione():
    """Restituisce la lista dei tipi di relazione."""
    tipi = get_tipi_relazione()
    return jsonify(tipi)


# ==============================================================================
# ROUTE: CERCA CLIENTI PER COLLEGAMENTO
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/cerca-clienti')
@login_required
def cerca_clienti_per_collegamento():
    """Cerca clienti da collegare (con filtro personale/globale)."""
    query = request.args.get('q', '').strip()
    cliente_id = request.args.get('cliente_id', type=int)
    solo_miei = request.args.get('solo_miei', 'false').lower() == 'true'
    
    if len(query) < 2:
        return jsonify([])
    
    conn = get_connection()
    cursor = conn.cursor()
    
    user_id = session.get('user_id')
    
    # Costruisci query base
    sql = '''
        SELECT id, COALESCE(ragione_sociale, nome_cliente) as ragione_sociale, p_iva, cod_fiscale, commerciale_id
        FROM clienti
        WHERE (ragione_sociale LIKE ? OR p_iva LIKE ? OR cod_fiscale LIKE ?)
        AND id != ?
    '''
    params = [f'%{query}%', f'%{query}%', f'%{query}%', cliente_id or 0]
    
    # Filtro "solo miei clienti"
    if solo_miei and user_id:
        subordinati = get_subordinati(conn, user_id)
        if subordinati:
            placeholders = ','.join('?' * len(subordinati))
            sql += f' AND commerciale_id IN ({placeholders})'
            params.extend(subordinati)
        else:
            conn.close()
            return jsonify([])
    
    sql += ' ORDER BY ragione_sociale LIMIT 20'
    
    cursor.execute(sql, params)
    columns = [desc[0] for desc in cursor.description]
    risultati = [dict(zip(columns, row)) for row in cursor.fetchall()]
    
    # Aggiungi identificativo
    for r in risultati:
        if r.get('p_iva'):
            r['identificativo'] = f"IT{r['p_iva']}"
        elif r.get('cod_fiscale'):
            r['identificativo'] = r['cod_fiscale']
        else:
            r['identificativo'] = None
    
    conn.close()
    return jsonify(risultati)


# ==============================================================================
# ROUTE: AGGIUNGI COLLEGAMENTO
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/aggiungi', methods=['POST'])
@login_required
def aggiungi_collegamento():
    """Crea un nuovo collegamento tra due clienti."""
    data = request.get_json()
    
    cliente_a_id = data.get('cliente_a_id')
    cliente_b_id = data.get('cliente_b_id')
    tipo_relazione = data.get('tipo_relazione', 'collegato')
    note = data.get('note', '').strip()
    
    if not cliente_a_id or not cliente_b_id:
        return jsonify({'success': False, 'error': 'ID clienti mancanti'}), 400
    
    if cliente_a_id == cliente_b_id:
        return jsonify({'success': False, 'error': 'Non puoi collegare un cliente a se stesso'}), 400
    
    user_id = session.get('user_id')
    
    conn = get_connection()
    cursor = conn.cursor()
    
    # Verifica se esiste gia' un collegamento attivo
    cursor.execute('''
        SELECT id FROM collegamenti_clienti
        WHERE ((cliente_a_id = ? AND cliente_b_id = ?) OR (cliente_a_id = ? AND cliente_b_id = ?))
        AND attivo = 1
    ''', (cliente_a_id, cliente_b_id, cliente_b_id, cliente_a_id))
    
    if cursor.fetchone():
        conn.close()
        return jsonify({'success': False, 'error': 'Collegamento gia\' esistente'}), 400
    
    # Inserisci collegamento
    cursor.execute('''
        INSERT INTO collegamenti_clienti (cliente_a_id, cliente_b_id, tipo_relazione, note, creato_da, creato_il)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (cliente_a_id, cliente_b_id, tipo_relazione, note or None, user_id, datetime.now()))
    
    conn.commit()
    collegamento_id = cursor.lastrowid
    conn.close()
    
    return jsonify({'success': True, 'id': collegamento_id})


# ==============================================================================
# ROUTE: RIMUOVI COLLEGAMENTO
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/rimuovi', methods=['POST'])
@login_required
def rimuovi_collegamento():
    """Disattiva un collegamento esistente."""
    data = request.get_json()
    collegamento_id = data.get('collegamento_id')
    
    if not collegamento_id:
        return jsonify({'success': False, 'error': 'ID collegamento mancante'}), 400
    
    user_id = session.get('user_id')
    
    conn = get_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        UPDATE collegamenti_clienti
        SET attivo = 0, disattivato_da = ?, disattivato_il = ?
        WHERE id = ?
    ''', (user_id, datetime.now(), collegamento_id))
    
    conn.commit()
    conn.close()
    
    return jsonify({'success': True})


# ==============================================================================
# ROUTE: STORICO COLLEGAMENTI (ADMIN)
# ==============================================================================

@collegamenti_bp.route('/admin/storico-collegamenti')
@login_required
@permesso_richiesto('gestione_utenti')
def storico_collegamenti():
    """Visualizza storico di tutti i collegamenti."""
    conn = get_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT 
            cc.id,
            cc.tipo_relazione,
            cc.note,
            cc.creato_il,
            cc.attivo,
            cc.disattivato_il,
            ca.ragione_sociale as cliente_a_nome,
            ca.p_iva as cliente_a_piva,
            cb.ragione_sociale as cliente_b_nome,
            cb.p_iva as cliente_b_piva,
            uc.nome || ' ' || uc.cognome as creato_da_nome,
            ud.nome || ' ' || ud.cognome as disattivato_da_nome
        FROM collegamenti_clienti cc
        JOIN clienti ca ON ca.id = cc.cliente_a_id
        JOIN clienti cb ON cb.id = cc.cliente_b_id
        LEFT JOIN utenti uc ON uc.id = cc.creato_da
        LEFT JOIN utenti ud ON ud.id = cc.disattivato_da
        ORDER BY cc.creato_il DESC
        LIMIT 500
    ''')
    
    columns = [desc[0] for desc in cursor.description]
    collegamenti = [dict(zip(columns, row)) for row in cursor.fetchall()]
    
    # Aggiungi descrizione relazione per lo storico
    for c in collegamenti:
        c['descrizione_relazione'] = get_descrizione_relazione(c['tipo_relazione'])
    
    conn.close()
    
    return render_template('admin/storico_collegamenti.html', collegamenti=collegamenti)


# ==============================================================================
# ROUTE: LISTA CLIENTI PER SELEZIONE GUIDATA
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/lista-clienti')
@login_required
def lista_clienti_collegamento():
    """Lista clienti per selezione guidata (con filtro personale/globale)."""
    cliente_id = request.args.get('cliente_id', type=int)
    solo_miei = request.args.get('solo_miei', 'false').lower() == 'true'
    
    conn = get_connection()
    cursor = conn.cursor()
    
    user_id = session.get('user_id')
    
    sql = '''
        SELECT id, COALESCE(ragione_sociale, nome_cliente) as ragione_sociale, p_iva, cod_fiscale, commerciale_id
        FROM clienti
        WHERE id != ?
    '''
    params = [cliente_id or 0]
    
    if solo_miei and user_id:
        subordinati = get_subordinati(conn, user_id)
        if subordinati:
            placeholders = ','.join('?' * len(subordinati))
            sql += f' AND commerciale_id IN ({placeholders})'
            params.extend(subordinati)
        else:
            conn.close()
            return jsonify([])
    
    sql += ' ORDER BY ragione_sociale LIMIT 500'
    
    cursor.execute(sql, params)
    columns = [desc[0] for desc in cursor.description]
    risultati = [dict(zip(columns, row)) for row in cursor.fetchall()]
    
    for r in risultati:
        if r.get('p_iva'):
            r['identificativo'] = f"IT{r['p_iva']}"
        elif r.get('cod_fiscale'):
            r['identificativo'] = r['cod_fiscale']
        else:
            r['identificativo'] = 'N/D'
    
    conn.close()
    return jsonify(risultati)


# ==============================================================================
# ROUTE: MODIFICA COLLEGAMENTO
# ==============================================================================

@collegamenti_bp.route('/api/collegamenti/modifica', methods=['POST'])
@login_required
def modifica_collegamento():
    """Modifica tipo relazione o note di un collegamento esistente."""
    data = request.get_json()
    
    collegamento_id = data.get('collegamento_id')
    tipo_relazione = data.get('tipo_relazione')
    note = data.get('note', '').strip()
    
    if not collegamento_id:
        return jsonify({'success': False, 'error': 'ID collegamento mancante'}), 400
    
    if not tipo_relazione:
        return jsonify({'success': False, 'error': 'Tipo relazione mancante'}), 400
    
    conn = get_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        UPDATE collegamenti_clienti
        SET tipo_relazione = ?, note = ?
        WHERE id = ? AND attivo = 1
    ''', (tipo_relazione, note or None, collegamento_id))
    
    conn.commit()
    conn.close()
    
    return jsonify({'success': True})
